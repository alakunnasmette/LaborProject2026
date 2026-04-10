# prognosis_model.py
import openpyxl
import os
import tkinter as tk
from tkinter import ttk
from prognosis_excel_mapping import process_submission
from ui.ui_components import add_nav_buttons, create_sidebar, add_logo_to_sidebar
from ui.ui_styles import COLOR_PRIMARY

_USE_PILLOW = True
EXCEL_PATH = "Integratie_Prognose_Model_5.0(1).xlsx"

ROW_BG_1 = "#EEEEEE"
ROW_BG_2 = "#E0E0E0"

# --------- Prognosis model page begin ---------

def build_prognosis_page(parent, client=None, go_back=None):
    print(f"build_prognosis_page aangeroepen: parent={parent}, client={client}, go_back={go_back}")
    parent.configure(bg="white")

    # --------- Hoofd frame ---------
    page = tk.Frame(parent, bg="white")
    page.pack(fill="both", expand=True)

    # --------- Sidebar ---------

    # --------- Terugknop --------
    def handle_back():
        print("KNOP GEKLIKT")  # debug
        print(f"handle_back aangeroepen, go_back={go_back}, client={client}")
        if go_back:
            for w in parent.winfo_children():
                w.destroy()
            go_back(client)            


    # --------- Page header ---------
    header = tk.Frame(page, bg="white")
    header.pack(fill="x", padx=30, pady=(20, 10))

    tk.Label(
        header,
        text="Integratie prognosemodel",
        bg="white",
        fg="black",
        font=("Segoe UI", 14, "bold"),
        anchor="w"
    ).pack(anchor="w")

    tk.Label(
        header,
        text="Deze vragenlijst brengt uw persoonlijke situatie, vaardigheden en arbeidsmarktpositie in kaart. Op basis van uw antwoorden wordt een prognose opgesteld over uw kansen op de arbeidsmarkt.",
        bg="white",
        fg="#000000",
        font=("Segoe UI", 10, "normal"),
        anchor="w",
        wraplength=900,
        justify="left"
    ).pack(anchor="w", pady=(4, 0))

    # ========== Scroll container ==========
    container = tk.Frame(page, bg="white")
    container.pack(fill="both", expand=True, padx=20, pady=10)

    canvas = tk.Canvas(container, bg="white", highlightthickness=0)
    canvas.pack(side="left", fill="both", expand=True)

    scrollbar = ttk.Scrollbar(container, orient="vertical", command=canvas.yview)
    scrollbar.pack(side="right", fill="y")

    canvas.configure(yscrollcommand=scrollbar.set)

    content = tk.Frame(canvas, bg="#FFFFFF")
    canvas_window = canvas.create_window((0, 0), window=content, anchor="nw")

    def on_canvas_configure(event):
        canvas.itemconfig(canvas_window, width=event.width)

    canvas.bind("<Configure>", on_canvas_configure)

    def on_configure(event):
        canvas.configure(scrollregion=canvas.bbox("all"))
        canvas.itemconfig(canvas_window, width=event.width)

    content.bind("<Configure>", on_configure)

    def on_mousewheel(event):
        canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

    canvas.bind_all("<MouseWheel>", on_mousewheel)

    # --------- Lijst voor alle antwoorden ---------
    all_vars = []

    # --------- Hulpfuncties ---------
    def make_question_row(parent, number: int, question: str, options: list):
        bg = ROW_BG_1 if number % 2 != 0 else ROW_BG_2

        row_frame = tk.Frame(parent, bg=bg)
        row_frame.pack(fill="x", pady=1, padx=0)

        tk.Label(
            row_frame,
            text=str(number),
            width=4,
            bg="#F1C40F",
            fg="#000000",
            font=("Segoe UI", 9, "bold"),
            pady=6,
            relief="flat"
        ).pack(side="left", padx=(8, 6), pady=4)

        text_frame = tk.Frame(row_frame, bg=bg)
        text_frame.pack(side="left", fill="x", expand=True, pady=4, padx=(0, 10))

        tk.Label(
            text_frame,
            text=question,
            bg=bg,
            fg="#111111",
            font=("Segoe UI", 9),
            anchor="w",
            justify="left",
            wraplength=750
        ).pack(anchor="w")

        vars_dict = {}
        opts_frame = tk.Frame(text_frame, bg=bg)
        opts_frame.pack(anchor="w", pady=(3, 2))

        for opt in options:
            var = tk.BooleanVar(value=False)
            vars_dict[opt] = var
            tk.Checkbutton(
                opts_frame,
                text=opt,
                variable=var,
                bg=bg,
                fg="#111111",
                selectcolor="#FFFFFF",
                activebackground=bg,
                activeforeground="#111111",
                font=("Segoe UI", 9),
                relief="flat",
                cursor="hand2",
            ).pack(side="left", padx=(0, 14))

        return vars_dict

    def make_section_header(text):
        tk.Label(
            content,
            text=text,
            bg="#F1C40F",
            fg="black",
            padx=10,
            pady=6,
            anchor="w",
            font=("Segoe UI", 9, "bold")
        ).pack(fill="x", pady=(10, 0))

    def make_column_header():
        bar = tk.Frame(content, bg="#E3B911")
        bar.pack(fill="x")
        tk.Label(bar, text="Nr.", width=4, bg="#E3B911", fg="black",
                 font=("Segoe UI", 8, "bold"), anchor="w", padx=8, pady=4).pack(side="left")
        tk.Label(bar, text="Vraag", bg="#E3B911", fg="black",
                 font=("Segoe UI", 8, "bold"), anchor="w", padx=2, pady=4).pack(side="left")

    # ===== VRAGEN =====

    # SECTIE 1
    make_section_header("1. Persoonlijk")
    make_column_header()
    all_vars.append(make_question_row(content, 1, "Bent u een man of vrouw?", ["Man", "Vrouw", "Anders"]))
    all_vars.append(make_question_row(content, 2, "Wat is uw leeftijd?", ["15-25", "25-45", "46-65", "> 65"]))
    all_vars.append(make_question_row(content, 3, "Wat is uw huidige nationaliteit?", ["Autochtoon", "Westers", "Niet-westers"]))
    all_vars.append(make_question_row(content, 4, "In welke provincie bent u opzoek naar werk?", ["Zeeland", "Utrecht", "Noord-Brabant", "Gelderland", "Noord-Holland", "Overijsel", "Limburg", "Drenthe", "Friesland", "Flevoland", "Zuid-Holland", "Groningen"]))
    all_vars.append(make_question_row(content, 5, "Wat is uw migratieachtergrond?", ["Migratieachtergrond: Autochtoon", "Migratieachtergrond: Westers", "Migratieachtergrond: Niet Westers"]))

    # SECTIE 2
    make_section_header("2. Persoonlijkheid & Representativiteit")
    make_column_header()
    all_vars.append(make_question_row(content, 6, "Wat zijn uw meeste dominante persoonlijkheidskenmerken? (beroeps- of loopbaanonderzoek)", ["Openheid", "Consciëntieusheid", "Extraversie", "Altruïstisch", "Neuroticisme"]))
    all_vars.append(make_question_row(content, 7, "Mate van representativiteit. (Beoordeling door consulent)", ["Informeel", "Gekleed Informeel", "Zakelijk Informeel", "Zakelijk Formeel", "Gekleed Formeel"]))

    # SECTIE 3
    make_section_header("3. Gezondheid & Vitaliteit")
    make_column_header()
    all_vars.append(make_question_row(content, 8, "Ervaart u psychische en/of lichamelijke problemen? Indien, welke problematiek ervaart u? (Mogelijke inzage in BA/VA/AD onderzoek)", ["Geen", "Psychisch", "Somatisch", "Psychisch & Somatisch"]))
    all_vars.append(make_question_row(content, 9, "Vormen uw gezondheidsbeperkingen een belemmering voor uw integratie? Indien, in welke mate bent u arbeidsgeschikt dan wel arbeidsongeschikt? (Mogelijke inzage in BA/VA/AD onderzoek)", ["Volledige en duurzaam arbeidsongeschikt", "Gedeeltelijk arbeidsgeschikt", "Arbeidsgehandicapt", "Personen met een minder dan goed ervaren gezondheid", "Geen gezondheidsbeperkingen"]))
    all_vars.append(make_question_row(content, 10, "Beoefent u een sport? Indien, beschrijf de activiteit en de frequentie waarin u deze activiteit beoefent.", ["Sedentair", "Recreatief", "Actief"]))

    # SECTIE 4
    make_section_header("4. Financiën & Uitkering")
    make_column_header()
    all_vars.append(make_question_row(content, 11, "Ontvangt u een uitkering? Indien, welke uitkering ontvangt u?", ["WW", "Bijstand", "Geen"]))
    all_vars.append(make_question_row(content, 12, "Hoelang ontvangt u reeds uw WW uitkering?", ["< 3 maanden", "3 - 6 maanden", "6 - 12 maanden", "> 12 - 18 maanden", "18 - 24 maanden"]))
    all_vars.append(make_question_row(content, 13, "Hoelang ontvangt u reeds uw Bijstandsuitkering?", ["< 1 jaar", "1 - 2 jaar", "2 - 3 jaar", "3 - 4 jaar", "4 - 5 jaar"]))
    all_vars.append(make_question_row(content, 14, "Ervaart u schuldenproblematiek? Indien, doorloopt u een schuldhulpverlenings- of een schuldsaneringstraject?", ["Schuldsanering", "Schuldhulpverlening", "Geen schuldenproblematiek"]))
    all_vars.append(make_question_row(content, 15, "Ontvangt u naast uw uitkering nog overige inkomsten? Bijvoorbeeld uit zelfstandig ondernemerschap of een bijbaan. Indien, hebben deze inkomsten invloed op uw uitkering?", ["Overige inkomsten zonder invloed op uitkering", "Overige inkomsten met invloed op uitkering", "Geen overige inkomsten"]))
    all_vars.append(make_question_row(content, 16, "Ontvangt u naast uw uitkering en mogelijke overige inkomsten ook financiële steun vanuit uw omgeving?", ["Beschikbaar", "Niet beschikbaar"]))

    # SECTIE 5
    make_section_header("5. Cognitie & Taal")
    make_column_header()
    all_vars.append(make_question_row(content, 17, "Wat is de hoogte van uw IQ?", ["Zwak begaafd (IQ 70-85)", "Normaal (IQ 85-115)", "Meer begaafd (IQ 115-130)", "Hoog begaafd (IQ > 130)"]))
    all_vars.append(make_question_row(content, 18, "Beschrijf uw mate van Nederlandse taalvaardigheid.", ["Basisgebruiker (A1-A2 / 20% van Nederlandse bevolking)", "Onafhankelijke gebruiker (B1-B2 / 65% van Nederlandse bevolking)", "Vaardige gebruiker (C1-C2 / 15% van Nederlandse bevolking)"]))
    all_vars.append(make_question_row(content, 19, "Spreekt u nog vreemde talen? Indien, welke talen spreekt u nog meer?", ["Geen", "1 vreemde taal", "2 of meer vreemde talen"]))
    all_vars.append(make_question_row(content, 20, "Beschrijf uw mate van buitenlandse taalvaardigheid.", ["Basisgebruiker (A1-A2)", "Onafhankelijke gebruiker (B1-B2)", "Vaardige gebruiker (C1-C2)"]))
    all_vars.append(make_question_row(content, 21, "Beschrijf uw ICT-competenties en het niveau waarop u deze kunt uitdragen.", ["Geen", "Basis", "Gevorderd (MBO 3-4)", "Expert (HBO/WO)"]))

    # SECTIE 6
    make_section_header("6. Onderwijs")
    make_column_header()
    all_vars.append(make_question_row(content, 22, "Wat is uw hoogst behaalde opleiding?", ["Onderwijsniveau: Laag", "Onderwijsniveau: Middelbaar", "Onderwijsniveau: Hoog"]))
    all_vars.append(make_question_row(content, 23, "Indien u MBO bent afgestudeerd, welke opleiding(en) heeft u gevolgd?", ["Matig-slechte perspectieven", "Redelijke perspectieven", "Goede perspectieven"]))
    all_vars.append(make_question_row(content, 24, "Indien u HBO bent afgestudeerd, welke opleiding(en) heeft u gevolgd?", ["Matig-slechte perspectieven", "Redelijke perspectieven", "Goede perspectieven"]))
    all_vars.append(make_question_row(content, 25, "Indien u WO bent afgestudeerd, welke opleiding(en) heeft u gevolgd?", ["Matig-slechte perspectieven", "Redelijke perspectieven", "Goede perspectieven"]))

    # SECTIE 7
    make_section_header("7. Persoonlijke Situatie")
    make_column_header()
    all_vars.append(make_question_row(content, 26, "Beschrijf uw huidige leefsituatie.", ["Alleenstaand", "Samenwonend/gehuwd zonder kind(eren)", "Samenwonend/gehuwd met kind(eren) <9 jaar", "Samenwonend/gehuwd met kind(eren) >9 jaar"]))
    all_vars.append(make_question_row(content, 27, "Wat is de hoogte van uw reguliere jaarinkomen?", ["Laagste inkomenklasse", "Laag midden inkomensklasse", "Midden inkomensklasse","Hoog midden inkomensklasse", "Hoogste inkomenklasse"]))
    all_vars.append(make_question_row(content, 28, "Wat is uw hoogst behaalde onderwijsniveau?", ["Laag onderwijsniveau(LBO)", "Middelbaar onderwijsniveau (MBO)", "Hoger onderwijsniveau (HBO/WO)"]))
    all_vars.append(make_question_row(content, 29, "Hoelang bevindt u zich reeds in het naturalisatieproces?", ["1-3 jaar", "4 jaar", "5 jaar", "6 jaar", "7 jaar", "8-10 jaar"]))
    all_vars.append(make_question_row(content, 30, "Hoe groot is de omvang van uw netwerk? Beschikt u over LinkedIn? Indien, beschrijf de frequentie en de intensiteit van uw gebruik.", ["Klein (< 25 pers.)", "Gemiddeld (25-50 pers.)", "Groot (> 50 pers.)"]))
    all_vars.append(make_question_row(content, 31, "Ervaart u problemen met justitie? Indien, vormen deze problemen een belemmering voor werk en kunt u een VOG overleggen?", ["Problemen met justitie | geen VOG overlegbaar", "Problemen met justitie | VOG overlegbaar", "Geen problemen met justitie"]))
    all_vars.append(make_question_row(content, 32, "Ervaart u problemen in uw thuissituatie? Indien, vormen deze problemen een belemmering voor werk?", ["Problemen in thuissituatie | Belemmering voor werk", "Problemen in thuissituatie | Geen belemmering voor werk", "Geen problemen in thuissituatie"]))
    all_vars.append(make_question_row(content, 33, "Ervaart u verslavingsproblematiek? Indien, vormt deze problematiek een belemmering voor werk?", ["Verslavingsproblematiek | Belemmering voor werk", "Verslavingsproblematiek | Geen belemmering voor werk", "Geen verslavingsproblematiek"]))
    all_vars.append(make_question_row(content, 34, "Bent u in het bezit van een rijbewijs?", ["Nee", "ik volg lessen", "Ja"]))
    all_vars.append(make_question_row(content, 35, "Beschikt u over eigen vervoer en bent u in staat om zelfstandig naar uw werkplek te reizen? Indien u afhankelijk bent van het OV, ervaart u dan beperkingen m.b.t. aansluitingen?", ["Beperkt(OV)", "Gedeeltelijk beperkt(OV)", "Onbeperkt(OV/EV)"]))

    # SECTIE 8
    make_section_header("8. Arbeidsmarkt & Re-integratie")
    make_column_header()
    all_vars.append(make_question_row(content, 36, "Beschrijf uw arbeidsmarkttransitie.", ["I: Betaalde arbeid", "II: Scholing naar betaalde arbeid", "III: Zorg/Huishouden naarbetaalde arbeid", "IV: Uittreding naar bet.Arbeid"]))
    all_vars.append(make_question_row(content, 37, "Beschrijf uw arbeidsmarkttransitie.", ["Vanuit volledige arbeidsongeschiktheid", "Vanuit gedeeltelijke arbeidsongeschiktheid", "Vanuit ziekte", "Vanuit ontslag/ Vanuit startpositie"]))
    all_vars.append(make_question_row(content, 38, "Wilt u solliciteren in een gelijkwaardige of een ongelijkwaardige sector als voorheen?", ["Homogene transitie", "Heterogene transitie"]))
    all_vars.append(make_question_row(content, 39, "Bent u geheel of slechts gedeeltelijk werkloos? Heeft u bijvoorbeeld een deeltijdbaan?", ["Volledig werkloos", "Gedeeltelijk werkloos"]))
    all_vars.append(make_question_row(content, 40, "Verricht u momenteel vrijwilligerswerk?", ["Vrijwilligerwerk", "Geen vrijwilligerswerk"]))
    all_vars.append(make_question_row(content, 41, "Hoelang bent u reeds werkloos?", ["0-3 mnd", "3-6 mnd", "6-12 mnd", "> 12 mnd"]))
    all_vars.append(make_question_row(content, 42, "Bent u op zoek naar structureel (vast) werk of seizoenswerk?", ["Structureel", "Seizoen"]))
    all_vars.append(make_question_row(content, 43, "Hoeveel jaar aan relevante werkervaring heeft u? (CV analyse)", ["1-5 jaar (Starter)", "6-10 jaar", "11-15 jaar", ">15 jaar"]))
    all_vars.append(make_question_row(content, 44, "Hoeveel sollicitaties verricht u per maand?", ["1-4", "5-10", ">10"]))
    all_vars.append(make_question_row(content, 45, "In welke periode solliciteert men momenteel? (Beoordeling door consulent)", ["Jan-Mrt", "Apr-Jun", "Jul-Sept", "Okt-Dec"]))

    # =========================================================================
    # Excel opslaan
    # =========================================================================
    def on_submit():
        answers = {}

        for i, vars_dict in enumerate(all_vars):
            options = list(vars_dict.keys())

        for j, opt in enumerate(options):
            if vars_dict[opt].get():
                q_key = f"q{i+1}"
                answers[q_key] = j
                break

    # 👇 NU BINNEN DE FUNCTIE
        if client:
            client_id = f"{client['id']}_{client['name'].replace(' ', '_')}"
        else:
            client_id = "onbekend"

    # 👇 HIER gebeurt alles
        output_path = process_submission(client_id, answers)

        print(f"Prognose opgeslagen: {output_path}")

        if go_back:
            for w in parent.winfo_children():
                w.destroy()
            go_back(client)

    # =========================================================================
    # Excel inladen
    # =========================================================================
    def load_from_excel():
        if not os.path.exists(EXCEL_PATH):
            return
        wb = openpyxl.load_workbook(EXCEL_PATH)
        ws = wb.active
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
            if i < len(all_vars) and row[1]:
                opgeslagen = [opt.strip() for opt in str(row[1]).split(",")]
                for opt, var in all_vars[i].items():
                    var.set(opt in opgeslagen)

    load_from_excel()

    # --------- Knoppen onderaan ---------
    btn_frame = tk.Frame(page, bg="white")
    btn_frame.pack(fill="x", padx=20, pady=10)

    add_nav_buttons(
        btn_frame,
        submit_command=on_submit,
        skip_command=handle_back,
        skip_text="Overslaan",
        submit_text="Opslaan en verder",
        skip_side="left",
        submit_side="right",
        padx=20,
    )

# --------- Standalone test ---------
if __name__ == "__main__":
    root = tk.Tk()
    root.geometry("800x600")
    build_prognosis_page(root)
    root.mainloop()

# --------- Pagina einde prognosis model ---------